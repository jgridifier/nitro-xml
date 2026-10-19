# SPDX-License-Identifier: MIT
"""Regenerates notes.docx and deck.pptx (original content; needs python-pptx + python-docx).

    python3 -I demo/samples/make_samples.py
"""
import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PC
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt as PPt

HERE = os.path.dirname(os.path.abspath(__file__))

# ---- docx ----
d = Document()
d.core_properties.title = "Quarterly field notes"
st = d.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
d.add_heading("Quarterly field notes", 0)
p = d.add_paragraph("Prepared for the product team · Q3 summary")
p.runs[0].italic = True
p.runs[0].font.color.rgb = RGBColor(0x60, 0x5E, 0x5C)
d.add_paragraph(
    "This sample document was written for the Nitro XML demo. It is rendered by a Rust and "
    "WebAssembly engine entirely in your browser: no server, no upload."
)
d.add_heading("What we shipped", 1)
for t in ("A faster first paint for large workbooks", "Embedded toolbars for all three formats", "Right-to-left text that lays out correctly"):
    d.add_paragraph(t, style="List Bullet")
d.add_heading("Numbers at a glance", 1)
rows = [("Metric", "Q2", "Q3"), ("Files opened", "12,400", "18,950"), ("Median open time", "410 ms", "290 ms"), ("Crash-free sessions", "99.2%", "99.7%")]
tb = d.add_table(rows=len(rows), cols=3)
tb.style = "Light Grid Accent 1"
for r, row in enumerate(rows):
    for c, v in enumerate(row):
        cell = tb.cell(r, c)
        cell.text = v
        if c:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
d.add_heading("Next steps", 1)
q = d.add_paragraph()
q.add_run("Bold").bold = True
q.add_run(", ")
q.add_run("italic").italic = True
q.add_run(" and ")
u = q.add_run("underlined")
u.underline = True
q.add_run(" runs, plus a numbered list:")
for t in ("Collect feedback from embedders", "Land the editing commands", "Publish the benchmark report"):
    d.add_paragraph(t, style="List Number")
d.save(os.path.join(HERE, "notes.docx"))

# ---- pptx ----
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]
ACC = PC(0xC4, 0x3E, 0x1C)
INK = PC(0x24, 0x24, 0x24)


def text(s, x, y, w, h, t, size, bold=False, color=INK, align=None):
    tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    tf.text = t
    r = tf.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = PPt(size), bold, color, "Calibri"
    if align:
        tf.paragraphs[0].alignment = align
    return tf


def box(s, kind, x, y, w, h, fill):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def bg(s, color):
    f = s.background.fill
    f.solid()
    f.fore_color.rgb = color


s = prs.slides.add_slide(blank)
bg(s, ACC)
text(s, 0.9, 2.3, 11.5, 1.6, "Nitro XML", 66, True, PC(255, 255, 255))
text(s, 0.9, 3.9, 11.5, 1, "Open Office files in the browser", 28, False, PC(0xFF, 0xE3, 0xD9))
box(s, MSO_SHAPE.RECTANGLE, 0.9, 3.75, 1.6, 0.06, PC(255, 255, 255))

s = prs.slides.add_slide(blank)
text(s, 0.8, 0.5, 11.7, 1, "Why it matters", 40, True, ACC)
for i, (h, b) in enumerate((("Private", "Files never leave the device"), ("Fast", "Rust and WebAssembly engine"), ("Familiar", "Looks and behaves like Office"))):
    x = 0.8 + i * 4.0
    box(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.0, 3.6, 3.8, PC(0xF3, 0xF2, 0xF1))
    box(s, MSO_SHAPE.OVAL, x + 0.3, 2.3, 0.8, 0.8, ACC)
    text(s, x + 0.3, 3.4, 3.0, 0.7, h, 26, True)
    text(s, x + 0.3, 4.1, 3.0, 1.4, b, 18, False, PC(0x60, 0x5E, 0x5C))

s = prs.slides.add_slide(blank)
text(s, 0.8, 0.5, 11.7, 1, "Open time by format", 40, True, ACC)
for i, (n, v, c) in enumerate((("xlsx", 5.2, PC(0x10, 0x7C, 0x41)), ("docx", 3.6, PC(0x18, 0x5A, 0xBD)), ("pptx", 4.4, ACC))):
    y = 2.0 + i * 1.5
    text(s, 0.8, y + 0.2, 1.4, 0.8, n, 24, True)
    box(s, MSO_SHAPE.RECTANGLE, 2.3, y, v * 1.9, 1.0, c)
    text(s, 2.4 + v * 1.9, y + 0.2, 2, 0.8, f"{v * 60:.0f} ms", 22)

s = prs.slides.add_slide(blank)
text(s, 0.8, 0.5, 11.7, 1, "Roadmap", 40, True, ACC)
tbl = s.shapes.add_table(4, 3, Inches(0.8), Inches(1.9), Inches(11.7), Inches(3.2)).table
for r, row in enumerate((("Quarter", "Focus", "Status"), ("Q3", "Viewer parity", "Shipped"), ("Q4", "Editing commands", "In progress"), ("Q1", "Collaboration", "Planned"))):
    for c, v in enumerate(row):
        tbl.cell(r, c).text = v
        tbl.cell(r, c).text_frame.paragraphs[0].runs[0].font.size = PPt(22)

s = prs.slides.add_slide(blank)
bg(s, PC(0x24, 0x24, 0x24))
text(s, 0.9, 2.8, 11.5, 1.4, "Thank you", 60, True, PC(255, 255, 255), 2)
text(s, 0.9, 4.3, 11.5, 0.8, "github.com/jgridifier/nitro-xml", 22, False, PC(0xBD, 0xBB, 0xB9), 2)
prs.save(os.path.join(HERE, "deck.pptx"))
